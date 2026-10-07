"""Disposable provider inputs only; native app execution/reopen remains separately verified."""
import hashlib
import json
import os
import sys
from copy import copy
from pathlib import Path

from PIL import Image
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from docx import Document
from docx.shared import Inches
from pptx import Presentation
from pptx.util import Inches as SlideInches

directory = Path(sys.argv[1]).resolve()
qualification = (Path(os.environ["LOCALAPPDATA"]) / "gotzji" / "qualification").resolve()
if qualification not in directory.parents or directory.exists():
    raise RuntimeError("NATIVE_FIXTURE_DIRECTORY_DENIED")
directory.mkdir(parents=True)
picture = directory / "preserved.png"
Image.new("RGB", (16, 16), (0, 128, 128)).save(picture)
book = Workbook()
sheet = book.active
sheet.title = "Fixture"
sheet["A1"] = "before"
sheet["B1"] = "=1+2"
sheet["C3"] = "preserved"
preserved_font = copy(sheet["C3"].font)
preserved_font.bold = True
sheet["C3"].font = preserved_font
sheet.add_image(ExcelImage(picture), "E5")
book.save(directory / "original.xlsx")
document = Document()
document.add_paragraph("before")
document.add_paragraph().add_run("preserved").bold = True
document.add_picture(str(picture), width=Inches(0.25))
document.save(directory / "original.docx")
presentation = Presentation()
slide = presentation.slides.add_slide(presentation.slide_layouts[6])
editable = slide.shapes.add_textbox(SlideInches(1), SlideInches(1), SlideInches(3), SlideInches(1))
editable.name = "editable"
editable.text_frame.text = "before"
preserved = slide.shapes.add_textbox(SlideInches(1), SlideInches(2), SlideInches(3), SlideInches(1))
preserved.name = "preserved"
preserved.text_frame.text = "preserved"
preserved.text_frame.paragraphs[0].runs[0].font.bold = True
slide.shapes.add_picture(str(picture), SlideInches(1), SlideInches(3), width=SlideInches(0.25))
presentation.save(directory / "original.pptx")
rows = [
    {"provider": "excel", "filePath": str(directory / "original.xlsx"), "sheet": "Fixture", "range": "A1"},
    {"provider": "word", "filePath": str(directory / "original.docx"), "paragraph": 1},
    {"provider": "powerpoint", "filePath": str(directory / "original.pptx"), "slide": 1, "shape": "editable"},
]
for row in rows:
    row["expectedSha256"] = hashlib.sha256(Path(row["filePath"]).read_bytes()).hexdigest()
print(json.dumps(rows))
