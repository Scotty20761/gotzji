"""Regenerate committed synthetic native fixtures; not a test-runtime dependency.
Authoring dependencies: python-docx 1.2.0 and python-pptx 1.0.2.
"""
from datetime import datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED
import base64
from docx import Document
from docx.shared import Inches as WordInches
from pptx import Presentation
from pptx.util import Inches as SlideInches, Pt

DESTINATION = Path(__file__).resolve().parent / 'native-office'
PNG = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR4nGMw7jjzHwAFAAKHH+5aBwAAAABJRU5ErkJggg==')


def stable_package(data: bytes, name: str) -> None:
    output = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(output, 'w', ZIP_DEFLATED) as target:
        for original in sorted(source.infolist(), key=lambda value: value.filename):
            entry = ZipInfo(original.filename, (2026, 10, 27, 0, 0, 0))
            entry.compress_type = ZIP_DEFLATED
            entry.external_attr = original.external_attr
            target.writestr(entry, source.read(original))
    DESTINATION.mkdir(parents=True, exist_ok=True)
    (DESTINATION / name).write_bytes(output.getvalue())


def properties(value) -> None:
    value.author = 'gotzji synthetic fixture'
    value.last_modified_by = 'gotzji synthetic fixture'
    value.title = 'Owned native preservation qualification'
    value.created = datetime(2026, 10, 27)
    value.modified = datetime(2026, 10, 27)


document = Document()
properties(document.core_properties)
document.add_paragraph('SELECTED_BEFORE')
document.add_paragraph().add_run('UNRELATED_KEEP').bold = True
document.add_table(rows=1, cols=1).cell(0, 0).text = 'TABLE_KEEP'
document.add_picture(BytesIO(PNG), width=WordInches(1))
word = BytesIO()
document.save(word)
stable_package(word.getvalue(), 'source.docx')

presentation = Presentation()
properties(presentation.core_properties)
slide = presentation.slides.add_slide(presentation.slide_layouts[6])
for name, text, top, size in [('SelectedText', 'SELECTED_BEFORE', 1, 18), ('PreservedText', 'UNRELATED_KEEP', 2, 14)]:
    box = slide.shapes.add_textbox(SlideInches(1), SlideInches(top), SlideInches(5), SlideInches(1))
    box.name = name
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.name = 'Arial'
    run.font.size = Pt(size)
slide.shapes.add_picture(BytesIO(PNG), SlideInches(1), SlideInches(3), width=SlideInches(1)).name = 'PreservedImage'
powerpoint = BytesIO()
presentation.save(powerpoint)
stable_package(powerpoint.getvalue(), 'source.pptx')